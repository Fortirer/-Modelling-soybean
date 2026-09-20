"""Minimal Markdown to Word converter, for the project report only.

Handles what FINAL_REPORT.md uses and nothing else: ATX headings, paragraphs,
bullet and numbered lists, pipe tables, images, block quotes, fenced code, and the
inline forms **bold**, *italic*, `code` and [text](url).

Why a hand-written converter. The docx skill's route is docx-js, and neither node
nor pandoc nor LibreOffice is installed on the machine this was built on, so
python-docx is used instead. The consequence to know about: the output could NOT
be rendered to check its layout. Structure is verified by reopening the file and
counting headings, tables and images; appearance is not.
"""
import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x1F, 0x2A, 0x37)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x55, 0x60, 0x6B)
HEADER_FILL = "DCE6F1"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\n]+\*|\[[^\]]+\]\([^)]+\))")


def _runs(par, text, size=None, color=None, italic=False, bold=False):
    """Add inline-formatted runs to a paragraph."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        r = None
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt((size or 10.5) - 1)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("[") and "](" in tok:
            label, url = tok[1:-1].split("](", 1)
            r = par.add_run(label)
            if url.startswith("http"):
                r2 = par.add_run(f" ({url})"); r2.font.size = Pt((size or 10.5) - 2)
                r2.font.color.rgb = MUTED
        else:
            r = par.add_run(tok)
        if r is not None:
            if size and not (tok.startswith("`")):
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
            if italic:
                r.italic = True
            if bold:
                r.bold = True


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    look = tblPr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(m)
    else:
        tblPr.append(m)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit"); e.set(qn("w:val"), "true"); trPr.append(e)


def _add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_margins(t)
    size = 9 if ncol <= 5 else (8 if ncol <= 8 else 7.5)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        _runs(c.paragraphs[0], h, size=size, bold=True)
        _shade(c, HEADER_FILL)
    _repeat_header(t.rows[0])
    for i, row in enumerate(body, start=1):
        _no_split(t.rows[i])
        for j in range(ncol):
            c = t.rows[i].cells[j]
            c.text = ""
            _runs(c.paragraphs[0], row[j] if j < len(row) else "", size=size)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def _page_number_footer(section, label):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label + "   |   Page ")
    r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    r2 = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r2._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
            r2._r.append(e)
    r2.font.size = Pt(8.5); r2.font.color.rgb = MUTED


def convert(md_text, out_path, base_dir, title_block, footer_label, contents=True):
    """Write `md_text` to `out_path`. Returns a dict of structural counts."""
    base = Path(base_dir)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    _page_number_footer(sec, footer_label)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12
    for name, sz, before, after in (("Heading 1", 17, 18, 6), ("Heading 2", 13.5, 12, 4),
                                    ("Heading 3", 11.5, 9, 3)):
        h = doc.styles[name]
        h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.bold = True
        h.font.color.rgb = ACCENT
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True

    # ---- title block --------------------------------------------------------------
    for text, sz, bold, color, after in title_block:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _runs(p, text, size=sz, bold=bold, color=color)
        p.paragraph_format.space_after = Pt(after)

    lines = md_text.split("\n")
    counts = dict(h1=0, h2=0, h3=0, tables=0, images=0, paragraphs=0, bullets=0)

    if contents:
        h1s = [re.sub(r"^#\s+", "", l) for l in lines if re.match(r"^##\s", l)]
        h1s = [re.sub(r"^##\s+", "", l) for l in lines if re.match(r"^##\s", l)]
        doc.add_heading("Contents", level=1)
        for h in h1s:
            p = doc.add_paragraph(style="List Bullet")
            _runs(p, h, size=10.5)
        doc.add_page_break()

    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("```"):                      # fenced code
            j = i + 1
            block = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                block.append(lines[j]); j += 1
            for b in block:
                p = doc.add_paragraph()
                r = p.add_run(b if b else " "); r.font.name = "Consolas"; r.font.size = Pt(8.5)
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(0)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            i = j + 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            text = m.group(2)
            if lvl == 1:                              # the document title is in the block
                i += 1
                continue
            lvl_doc = {2: 1, 3: 2, 4: 3}[lvl]
            h = doc.add_heading(level=lvl_doc)
            _runs(h, text)
            counts[f"h{lvl_doc}"] += 1
            i += 1
            continue
        if s.startswith("|"):                         # table
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            rows = [_split_row(b) for b in block if not re.match(r"^\|?\s*:?-{2,}", b.strip().strip("|").strip())]
            rows = [r for r in rows if any(c for c in r)]
            if len(rows) >= 1:
                _add_table(doc, rows); counts["tables"] += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", s)
        if m:
            alt, path = m.group(1), m.group(2)
            fp = base / path
            if fp.exists():
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(fp), width=Inches(6.4))
                p.paragraph_format.keep_with_next = True
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _runs(cap, alt, size=9, color=MUTED, italic=True)
                counts["images"] += 1
            else:
                p = doc.add_paragraph(); _runs(p, f"[missing figure: {path}]", color=MUTED)
            i += 1
            continue
        if s.startswith(">"):                         # block quote
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(re.sub(r"^>\s?", "", lines[i].strip())); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            _runs(p, " ".join(block), size=10, color=MUTED, italic=True)
            continue
        if re.match(r"^[-*]\s+", s):                  # bullets
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].strip()):
                p = doc.add_paragraph(style="List Bullet")
                _runs(p, re.sub(r"^[-*]\s+", "", lines[i].strip())); counts["bullets"] += 1
                i += 1
            continue
        if re.match(r"^\d+\.\s+", s):                 # numbered
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                p = doc.add_paragraph(style="List Number")
                _runs(p, re.sub(r"^\d+\.\s+", "", lines[i].strip())); counts["bullets"] += 1
                i += 1
            continue
        block = [s]                                   # paragraph
        i += 1
        while i < len(lines):
            n = lines[i].strip()
            if (not n or n == "---" or n.startswith(("#", "|", ">", "```", "![")) or
                    re.match(r"^[-*]\s+", n) or re.match(r"^\d+\.\s+", n)):
                break
            block.append(n); i += 1
        p = doc.add_paragraph()
        _runs(p, " ".join(block)); counts["paragraphs"] += 1

    z = doc.settings.element.find(qn("w:zoom"))
    if z is not None and z.get(qn("w:percent")) is None:
        z.set(qn("w:percent"), "100")
    doc.save(str(out_path))
    return counts
